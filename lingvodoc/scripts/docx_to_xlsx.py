import collections
import docx
import sys
import os
import re
import pyramid.paster as paster
import logging
import getopt
import xlsxwriter

# Setting up logging, if we are not being run as a script.
if __name__ != '__main__':
    log = logging.getLogger(__name__)
    log.debug('module init')


class Docx2XlsxError(Exception):

    def __init__(self, message):
        super().__init__(message)


def get_entries(filename, debug_flag=False):
    part = None
    is_left = False
    bold_words = ''
    left_words = ''
    right_words = ''
    result_dict = collections.defaultdict(list)
    header = 'n/a'
    total = 0

    log.debug(
        '\ndocx_path: {0}'.format(filename))

    try:
        doc = docx.Document(filename)

    except docx.opc.exceptions.PackageNotFoundError:
        raise Docx2XlsxError('input file is not a .docx format file')

    def write_out():
        nonlocal \
            bold_words, \
            left_words, \
            right_words, \
            result_dict, \
            header, \
            total

        if left_words:
            result_dict[header].append(
                {
                    'bold': bold_words,
                    'left': left_words,
                    'right': right_words
                }
            )

            if debug_flag:
                log.info(f'\nBold: {bold_words}\n'
                         f'Left: {left_words}\n'
                         f'Right: {right_words}\n'
                         f'---')

            bold_words = ''
            left_words = ''
            right_words = ''
            total += 1

    for p, para in enumerate(doc.paragraphs):
        # Start parsing on text like '- A -'
        if hdr := re.search(r'\u2014\s*\w\s*\u2014', para.text):
            write_out()
            header = hdr.group()
            if debug_flag:
                log.info(f'\nHeader: {header}\n')
            part = 'header'
            continue

        # Pause parsing if two lines in a row
        # are empty after dictionary body
        if (part == 'body' and
                not left_words and
                not para.text):
            part = None

        if not part or not para.text:
            continue

        # Checking if first word is bold
        first_is_bold = para.runs and (para.runs[0].bold or para.runs[0].font.cs_bold)

        if first_is_bold:
            part = 'body'
            is_left = True

        # Bold words or empty string mean that the previous dictionary
        # entry is over. Writing out collected words and clean variables.
        if first_is_bold or not para.text:
            write_out()

        # Getting bold text at beginning of string
        for word in para.runs:
            if not (word.bold or word.font.cs_bold):
                break
            bold_words += word.text

        for w, word in enumerate(para.runs):
            if is_left and ('\u2013 ' not in word.text):
                left_words += word.text
            elif not is_left:
                if w == 0:
                    right_words += f'\n{word.text.lstrip()}'
                else:
                    right_words += word.text
            else:
                dash_pos = word.text.index('\u2013 ')
                left_words += word.text[:dash_pos].rstrip()
                right_words += word.text[dash_pos+1:].lstrip()
                is_left = False

    write_out()

    if debug_flag:
        log.info(f'\nTotal dictionary entries: {total}')

    return result_dict


def write_xlsx(result_dict, xlsx_path, base_name, alpha_groups=False):

    def init_worksheet(worksheet):
        bold = workbook.add_format({'bold': True})
        worksheet.set_column(0, 0, 100)
        worksheet.set_column(1, 1, 200)
        worksheet.write('A1', '  Entry', bold)
        worksheet.write('B1', '  Meaning', bold)

    workbook = xlsxwriter.Workbook(xlsx_path)
    align = workbook.add_format()
    align.set_align('vcenter')
    align.set_text_wrap()

    if not alpha_groups:
        worksheet = workbook.add_worksheet(base_name)
        init_worksheet(worksheet)
        row = 1

    for header, group in result_dict.items():
        if alpha_groups:
            worksheet = workbook.add_worksheet(header)
            init_worksheet(worksheet)
            row = 1

        for line in group:
            h1 = (len(line['right']) // 200 + 1) * 17
            h2 = len(line['right'].splitlines()) * 17
            worksheet.set_row(row, max(h1, h2))
            for column, item in enumerate([line['left'], line['right']]):
                worksheet.write(row, column, item, align)
            row += 1

    workbook.close()

def main_import(args):

    # Parsing command-line options.
    opt_list, arg_list = (
        getopt.gnu_getopt(args, '', [
            'docx-file=',
            'xlsx-file=',
            'alpha-groups',
            'debug'
        ]))

    opt_dict = dict(opt_list)
    debug_flag = '--debug' in opt_dict
    alpha_groups = '--alpha-groups' in opt_dict

    if not (docx_path := opt_dict.get('--docx-file')):
        docx_path = arg_list[0]
    if not os.path.isfile(docx_path):
        raise FileNotFoundError(f'\nSpecified path {docx_path} is not correct.')

    if not (xlsx_path := opt_dict.get('--xlsx-file')):
        if len(arg_list) > 1:
            xlsx_path = arg_list[1]

    base_path = os.path.splitext(docx_path)[0]
    base_name = os.path.basename(base_path)

    if xlsx_path:
        os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)
    else:
        xlsx_path = base_path + '.xlsx'
        log.warning(f'\nSetting output path automatically: {os.path.abspath(xlsx_path)}')

    if os.path.isfile(docx_path):
        result_dict = get_entries(docx_path, debug_flag)
        write_xlsx(result_dict, xlsx_path, base_name, alpha_groups)
    else:
        log.error(
            f'\nSpecified path {docx_path} is not correct.')


if __name__ == "__main__":
    if (len(sys.argv) > 1 and
            sys.argv[1] == '-config'):

        # We have a configuration file; initializing logging.
        config_path = sys.argv[2]

        paster.setup_logging(config_path)
        log = logging.getLogger(__name__)

        arg_list = sys.argv[3:]
    else:
        # No config file, so just logging to stdout.

        arg_list = sys.argv[1:]

        log_root = logging.getLogger()
        log_root.setLevel(logging.DEBUG)

        log_handler = logging.StreamHandler(sys.stdout)
        log_handler.setLevel(logging.DEBUG)

        log_formatter = (
            logging.Formatter(
                '%(asctime)s %(levelname)-5.5s [%(name)s][%(threadName)s] '
                '%(pathname)s:%(lineno)d: %(message)s'))

        log_handler.setFormatter(log_formatter)
        log_root.addHandler(log_handler)

        log = logging.getLogger(__name__)

    # Doing what we need.

    if len(arg_list) <= 0:
        log.info(
            '\nPlease specify arguments.')
    else:
        main_import(arg_list)
